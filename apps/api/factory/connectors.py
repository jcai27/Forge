"""Reusable, permission-scoped business connectors; no credentials in manifests."""

import asyncio
import csv
import io
import json
import math
import socket
import ssl
import zlib
from abc import ABC, abstractmethod
from copy import deepcopy
from email.message import EmailMessage
from urllib.parse import urlparse

from fastapi import HTTPException
from sqlalchemy import select

from factory.http_transport import pinned_client, pinned_mcp_client, public_addresses, validate_url
from factory.models import Connection, DataSource, ToolDefinition
from factory.security import decrypt, redact_tenant


def safe_url(url, allowed_hosts):
    validate_url(url, allowed_hosts)
    parsed = urlparse(url)
    addresses = socket.getaddrinfo(
        parsed.hostname, parsed.port or (443 if parsed.scheme == "https" else 80), type=socket.SOCK_STREAM
    )
    public_addresses(addresses)
    return url


def tool(name, description, input_schema=None, output_schema=None, write=False):
    return {
        "name": name,
        "description": description,
        "input_schema": input_schema or {"type": "object", "properties": {}, "additionalProperties": False},
        "output_schema": output_schema or {"type": "object"},
        "write": write,
    }


class Connector(ABC):
    def __init__(self, config, secret="", records=None):
        self.config = config
        self.secret = secret
        self.records = records or {}

    @abstractmethod
    async def discover(self): ...

    @abstractmethod
    async def invoke(self, name, arguments): ...

    async def validate(self):
        resources = await self.discover()
        return {"ok": True, "resources": resources, "tools": self.tools()}

    def tools(self):
        return []


class ReplayConnector:
    """An ordered recording with exact argument matching; never opens a connection."""

    def __init__(self, responses):
        self.responses = deepcopy(responses)
        self.position = 0

    async def invoke(self, tool_id, arguments):
        if self.position >= len(self.responses):
            raise ValueError("Replay recording has no response for this call")
        expected = self.responses[self.position]
        if expected["tool_id"] != tool_id or expected["arguments"] != arguments:
            raise ValueError("Replay tool sequence or arguments differ from the recording")
        self.position += 1
        return deepcopy(expected["response"])


class FixtureConnector(Connector):
    async def discover(self):
        from factory.fixtures import fictional_data

        return [
            {"name": name, "schema": {"type": "array", "items": {"type": "object"}}, "sample": rows[:3]}
            for name, rows in fictional_data().items()
        ]

    def tools(self):
        return [
            tool(
                "prepare_operation",
                "Prepare a structured business operation using catalog, customer and deterministic policies. Never commits a write.",
                {"type": "object", "properties": {"input": {"type": "object"}}, "required": ["input"]},
            ),
            tool(
                "read_records",
                "Read an approved fictional business collection.",
                {"type": "object", "properties": {"collection": {"type": "string"}}, "required": ["collection"]},
            ),
        ]

    async def invoke(self, name, arguments):
        if name == "read_records":
            return {"records": self.records.get(arguments["collection"], [])}
        raise ValueError("prepare_operation is bound with versioned runtime context")


REST_REQUEST_TIMEOUT = 15


async def bounded_json(response, limit=2_000_000):
    """Bound both wire bytes and decoded bytes before parsing an upstream payload."""
    length = response.headers.get("Content-Length")
    if length is not None:
        try:
            advertised = int(length)
        except ValueError:
            raise ValueError("Invalid response length") from None
        if advertised < 0 or advertised > limit:
            raise ValueError("Response exceeds resource budget")
    encoding = response.headers.get("Content-Encoding", "identity").strip().lower()
    if encoding not in ("identity", "gzip", "deflate"):
        raise ValueError("Unsupported response encoding")
    decoder = (
        zlib.decompressobj(16 + zlib.MAX_WBITS if encoding == "gzip" else zlib.MAX_WBITS)
        if encoding != "identity"
        else None
    )
    body, received = bytearray(), 0
    chunks = response.aiter_raw(chunk_size=65536) if decoder else response.aiter_bytes(chunk_size=65536)
    async for chunk in chunks:
        received += len(chunk)
        if received > limit:
            raise ValueError("Response exceeds resource budget")
        if decoder:
            try:
                chunk = decoder.decompress(chunk, limit - len(body) + 1)
            except zlib.error:
                raise ValueError("Invalid compressed response") from None
        if len(body) + len(chunk) > limit:
            raise ValueError("Decoded response exceeds resource budget")
        body.extend(chunk)
    if decoder and (not decoder.eof or decoder.unused_data):
        raise ValueError("Incomplete or trailing compressed response")
    return json.loads(body)


class RESTConnector(Connector):
    def tools(self):
        return [
            tool(
                item["name"],
                item["description"],
                item.get("input_schema"),
                item.get("output_schema"),
                item.get("method", "GET").upper() not in ("GET", "HEAD"),
            )
            for item in self.config.get("operations", [])
        ]

    async def request(self, operation, arguments):
        base = validate_url(self.config["base_url"], self.config.get("allowed_hosts", []))
        path = operation["path"]
        if not path.startswith("/") or path.startswith("//") or ".." in path or "?" in path or "#" in path:
            raise ValueError("Operation path must be static and relative")
        url = base.rstrip("/") + path
        headers = {"Accept-Encoding": "gzip, deflate, identity"}
        if self.secret:
            headers["Authorization"] = f"Bearer {self.secret}"
        method = operation.get("method", "GET").upper()
        attempts = 3 if method in ("GET", "HEAD") else 1
        async with (
            asyncio.timeout(REST_REQUEST_TIMEOUT),
            pinned_client(base, self.config.get("allowed_hosts", []), timeout=15) as client,
        ):
            for attempt in range(attempts):
                async with client.stream(
                    method,
                    url,
                    headers=headers,
                    params=arguments if method in ("GET", "HEAD") else None,
                    json=arguments if method not in ("GET", "HEAD") else None,
                ) as response:
                    if response.status_code == 429 and attempt < attempts - 1:
                        try:
                            delay = float(response.headers.get("Retry-After", 1))
                        except ValueError:
                            delay = 1
                        delay = min(delay, 5) if math.isfinite(delay) and delay >= 0 else 1
                    else:
                        response.raise_for_status()
                        return await bounded_json(response)
                # Close discarded responses before waiting; writes never retry invisibly.
                await asyncio.sleep(delay)
        raise ValueError("Rate limit retries exhausted")

    async def discover(self):
        # Resource schemas must be explicitly supplied, never invented by the model.
        return self.config.get("resources", [])

    async def validate(self):
        operation = {"path": self.config.get("health_path", "/"), "method": "GET"}
        await self.request(operation, {})
        return await super().validate()

    async def invoke(self, name, arguments):
        operation = next((o for o in self.config.get("operations", []) if o["name"] == name), None)
        if not operation:
            raise ValueError("Unregistered REST operation")
        return await self.request(operation, arguments)


class PostgreSQLConnector(Connector):
    def connect(self):
        import psycopg

        # Store the entire DSN only in the encrypted secret; enforce read-only transactions.
        return psycopg.connect(
            self.secret, connect_timeout=10, options="-c default_transaction_read_only=on -c statement_timeout=10000"
        )

    async def discover(self):
        def read():
            with self.connect() as conn:
                rows = conn.execute(
                    "SELECT table_name, column_name, data_type FROM information_schema.columns WHERE table_schema = %s ORDER BY table_name, ordinal_position",
                    (self.config.get("schema", "public"),),
                ).fetchall()
            allowed = self.config.get("tables", [])
            tables = {}
            for table, column, type_name in rows:
                if table in allowed:
                    tables.setdefault(table, {})[column] = {"database_type": type_name}
            return [{"name": name, "schema": {"type": "object", "properties": cols}} for name, cols in tables.items()]

        return await asyncio.to_thread(read)

    def tools(self):
        return [
            tool(
                "read_table",
                "Read up to 100 records from an explicitly approved table. No arbitrary SQL.",
                {
                    "type": "object",
                    "properties": {
                        "table": {"type": "string", "enum": self.config.get("tables", [])},
                        "limit": {"type": "integer", "minimum": 1, "maximum": 100},
                    },
                    "required": ["table"],
                },
            )
        ]

    async def invoke(self, name, arguments):
        from psycopg import sql
        from psycopg.rows import dict_row

        if name != "read_table" or arguments["table"] not in self.config.get("tables", []):
            raise ValueError("Table access denied")
        limit = min(max(int(arguments.get("limit", 25)), 1), 100)

        def read():
            with self.connect() as conn:
                conn.row_factory = dict_row
                rows = conn.execute(
                    sql.SQL("SELECT * FROM {}.{} LIMIT %s").format(
                        sql.Identifier(self.config.get("schema", "public")), sql.Identifier(arguments["table"])
                    ),
                    (limit,),
                ).fetchall()
                return {"records": json.loads(json.dumps(rows, default=str))}

        return await asyncio.to_thread(read)


class FileConnector(Connector):
    @staticmethod
    def parse(content, filename):
        if len(content) > 5_000_000:
            raise ValueError("Files must be smaller than 5 MB")
        if filename.lower().endswith(".csv"):
            records = list(csv.DictReader(io.StringIO(content.decode("utf-8-sig"))))
        elif filename.lower().endswith(".xlsx"):
            from openpyxl import load_workbook

            workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
            sheet = workbook.active
            rows = sheet.iter_rows(values_only=True)
            headers = [str(cell) for cell in next(rows)]
            records = [dict(zip(headers, row, strict=False)) for row in rows]
            workbook.close()
        else:
            raise ValueError("Only CSV and XLSX are supported")
        if len(records) > 10000:
            raise ValueError("Maximum 10000 imported records")
        return json.loads(json.dumps(records, default=str))

    async def discover(self):
        return [
            {
                "name": name,
                "schema": {"type": "object", "properties": {k: {"type": "string"} for k in (rows[0] if rows else {})}},
                "sample": rows[:3],
            }
            for name, rows in self.records.items()
        ]

    def tools(self):
        return [
            tool(
                "read_file",
                "Read an explicitly approved imported resource.",
                {"type": "object", "properties": {"resource": {"type": "string"}}, "required": ["resource"]},
            )
        ]

    async def invoke(self, name, arguments):
        if name != "read_file" or arguments["resource"] not in self.records:
            raise ValueError("Resource is unavailable or not model-visible")
        return {"records": self.records[arguments["resource"]][:100]}


class WebhookConnector(Connector):
    async def discover(self):
        return [{"name": "inbound_events", "schema": self.config.get("input_schema", {"type": "object"})}]

    def tools(self):
        return [tool("read_events", "Read authenticated inbound webhook events.")]

    async def invoke(self, name, arguments):
        if name != "read_events":
            raise ValueError("Unregistered webhook tool")
        return {"records": self.records.get("inbound_events", [])[-100:]}


class EmailConnector(Connector):
    async def discover(self):
        return [
            {
                "name": "inbox",
                "schema": {"type": "object", "properties": {"subject": {"type": "string"}, "body": {"type": "string"}}},
            }
        ]

    def tools(self):
        return [
            tool("read_inbox", "Read recent email as untrusted data."),
            tool(
                "send_email",
                "Send email only after explicit human approval.",
                {
                    "type": "object",
                    "properties": {"to": {"type": "string"}, "subject": {"type": "string"}, "body": {"type": "string"}},
                    "required": ["to", "subject", "body"],
                },
                write=True,
            ),
        ]

    async def invoke(self, name, arguments):
        if self.config.get("provider") == "mock":
            if name == "read_inbox":
                return {
                    "messages": self.config.get(
                        "fixtures", [{"subject": "Order", "body": "Harbor Bistro: 2 kg Tomatoes"}]
                    )
                }
            if name == "send_email":
                return {"accepted": True, "mode": "mock", "recipient": arguments["to"]}

        def execute():
            credentials = json.loads(self.secret)
            if name == "read_inbox":
                import email
                import imaplib

                with imaplib.IMAP4_SSL(self.config["imap_host"], timeout=15) as client:
                    client.login(credentials["username"], credentials["password"])
                    client.select("INBOX", readonly=True)
                    status, ids = client.search(None, "ALL")
                    if status != "OK":
                        raise ValueError("IMAP search failed")
                    messages = []
                    for message_id in ids[0].split()[-10:]:
                        status, data = client.fetch(message_id, "(BODY.PEEK[])")
                        if status != "OK":
                            raise ValueError("IMAP fetch failed")
                        message = email.message_from_bytes(data[0][1])
                        parts = [part for part in message.walk() if part.get_content_type() == "text/plain"]
                        messages.append(
                            {
                                "subject": message.get("subject", ""),
                                "body": "\n".join(
                                    (part.get_payload(decode=True) or b"").decode(errors="replace")[:10000]
                                    for part in parts
                                ),
                            }
                        )
                    return {"messages": messages}
            if name == "send_email":
                import smtplib

                message = EmailMessage()
                message["From"] = credentials["username"]
                message["To"] = arguments["to"]
                message["Subject"] = arguments["subject"]
                message.set_content(arguments["body"])
                with smtplib.SMTP_SSL(
                    self.config["smtp_host"], context=ssl.create_default_context(), timeout=15
                ) as client:
                    client.login(credentials["username"], credentials["password"])
                    client.send_message(message)
                return {"accepted": True}
            raise ValueError("Unknown email operation")

        return await asyncio.to_thread(execute)


MCP_REQUEST_TIMEOUT = 15


class MCPConnector(Connector):
    def server(self):
        from agents.mcp import MCPServerStreamableHttp

        url = validate_url(self.config["url"], self.config.get("allowed_hosts", []))
        return MCPServerStreamableHttp(
            params={
                "url": url,
                "timeout": MCP_REQUEST_TIMEOUT,
                "sse_read_timeout": MCP_REQUEST_TIMEOUT,
                "headers": {"Authorization": f"Bearer {self.secret}"} if self.secret else {},
                "httpx_client_factory": lambda **kwargs: pinned_mcp_client(
                    url, self.config.get("allowed_hosts", []), **kwargs
                ),
            },
            client_session_timeout_seconds=MCP_REQUEST_TIMEOUT,
            max_list_pages=10,
            max_retry_attempts=0,
        )

    async def discover(self):
        async with asyncio.timeout(MCP_REQUEST_TIMEOUT), self.server() as server:
            tools = await server.list_tools()
            return [
                {
                    "name": t.name,
                    "schema": t.input_schema,
                    "description": t.description,
                    "write": t.name not in self.config.get("read_tools", []),
                }
                for t in tools
                if t.name in self.config.get("allowed_tools", [])
            ]

    def tools(self):
        return self.config.get("discovered_tools", [])

    async def invoke(self, name, arguments):
        if name not in self.config.get("allowed_tools", []):
            raise ValueError("MCP capability not approved")
        async with asyncio.timeout(MCP_REQUEST_TIMEOUT), self.server() as server:
            result = await server.call_tool(name, arguments)
            return result.model_dump(mode="json")


CONNECTORS = {
    "fixture": FixtureConnector,
    "rest": RESTConnector,
    "postgres": PostgreSQLConnector,
    "file": FileConnector,
    "webhook": WebhookConnector,
    "email": EmailConnector,
    "mcp": MCPConnector,
}


def adapter(db, connection, for_model=False, state=None):
    records = {}
    for source in db.scalars(
        select(DataSource).where(DataSource.connection_id == connection.id, DataSource.org_id == connection.org_id)
    ):
        if not for_model or source.model_visible:
            records[source.name] = source.records
    return CONNECTORS[connection.kind](
        connection.config, decrypt(connection.encrypted_secret), state if connection.kind == "fixture" else records
    )


def check_connection(db, org_id, connection_id, write=False):
    connection = db.scalar(
        select(Connection)
        .where(Connection.id == connection_id, Connection.org_id == org_id)
        .execution_options(populate_existing=True)
    )
    if not connection or connection.org_id != org_id or connection.status != "connected":
        raise HTTPException(403, "Connection unavailable or unauthorized")
    if not connection.permissions.get("write" if write else "read", False):
        raise HTTPException(403, "Connection permission denied")
    return connection


async def discover_connection(db, connection):
    result = await adapter(db, connection).validate()
    if connection.kind == "mcp":
        definitions = [
            tool(t["name"], t.get("description", ""), t["schema"], write=t["write"]) for t in result["resources"]
        ]
        connection.config = {**connection.config, "discovered_tools": definitions}
        result["tools"] = definitions
    result = redact_tenant(db, connection.org_id, result)
    existing = {
        t.name: t for t in db.scalars(select(ToolDefinition).where(ToolDefinition.connection_id == connection.id))
    }
    for definition in result["tools"]:
        row = existing.get(definition["name"])
        if row:
            row.definition = definition
            row.write = definition["write"]
        else:
            db.add(
                ToolDefinition(
                    org_id=connection.org_id,
                    connection_id=connection.id,
                    name=definition["name"],
                    definition=definition,
                    write=definition["write"],
                )
            )
    connection.status = "connected"
    return redact_tenant(db, connection.org_id, result)
